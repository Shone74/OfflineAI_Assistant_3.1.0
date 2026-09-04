"""Document loader — parses text/markdown/pdf/docx/html into chunks.

Optional backends (installed separately) enable real parsing of binary
formats; otherwise the loader degrades gracefully to raw-text extraction.
"""

from __future__ import annotations

import re
from pathlib import Path

from knowledge.models import Document, KnowledgeChunk


class DocumentLoader:
    """Loads documents and splits them into :class:`KnowledgeChunk` pieces."""

    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 100) -> None:
        self._chunk_size = chunk_size
        self._chunk_overlap = chunk_overlap

    def load(self, path: str | Path) -> Document:
        p = Path(path)
        doc = Document(path=str(p), title=p.stem, format=p.suffix.lstrip("."))
        doc.chunks = self._chunk_text(self._extract_text(p), doc.id)
        return doc

    # ------------------------------------------------------------------ #
    def _extract_text(self, path: Path) -> str:
        ext = path.suffix.lower()

        if ext == ".pdf":
            try:
                from pypdf import PdfReader

                return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
            except Exception:
                pass
            try:
                from fitz import open as fitz_open

                return "\n".join(pg.get_text() for pg in fitz_open(str(path)))
            except Exception:
                pass

        if ext == ".docx":
            try:
                from docx import Document as DocxDocument

                return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
            except Exception:
                pass
            try:
                from docx2txt import process

                return process(str(path))
            except Exception:
                pass

        if ext in (".html", ".htm"):
            return self._strip_html(path.read_text(encoding="utf-8", errors="replace"))

        return path.read_text(encoding="utf-8", errors="replace")

    @staticmethod
    def _strip_html(html: str) -> str:
        try:
            from bs4 import BeautifulSoup

            return BeautifulSoup(html, "html.parser").get_text("\n")
        except Exception:
            return re.sub(r"<[^>]+>", "", html)

    # ------------------------------------------------------------------ #
    def _chunk_text(self, text: str, doc_id: str) -> list[KnowledgeChunk]:
        text = (text or "").strip()
        if not text:
            return []
        chunks: list[KnowledgeChunk] = []
        step = self._chunk_size - self._chunk_overlap
        for start in range(0, len(text), max(step, 1)):
            chunks.append(
                KnowledgeChunk(doc_id=doc_id, text=text[start : start + self._chunk_size])
            )
        return chunks
